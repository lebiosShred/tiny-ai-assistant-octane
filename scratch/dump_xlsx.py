import openpyxl
import json

wb = openpyxl.load_workbook(r'C:\Users\SkyDr\Downloads\customer profiles.xlsx')
data = {}
for name in wb.sheetnames:
    sheet = wb[name]
    rows = []
    headers = [cell.value for cell in sheet[1]]
    headers = [h for h in headers if h is not None]
    
    for r in range(2, sheet.max_row + 1):
        row_values = [sheet.cell(r, c).value for c in range(1, len(headers) + 1)]
        if any(v is not None for v in row_values):
            row_dict = {}
            for i, h in enumerate(headers):
                if i < len(row_values):
                    row_dict[h] = row_values[i]
            rows.append(row_dict)
    data[name] = rows

with open('scratch/customer_profiles_dump.json', 'w', encoding='utf-8') as f:
    json.dump(data, f, indent=2, ensure_ascii=False)

print("Done")
