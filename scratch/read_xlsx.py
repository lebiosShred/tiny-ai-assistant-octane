import openpyxl

wb = openpyxl.load_workbook(r'C:\Users\SkyDr\Downloads\customer profiles.xlsx')
for name in wb.sheetnames:
    print(f"\n================ SHEET: {name} ================")
    sheet = wb[name]
    # Print headers
    headers = [cell.value for cell in sheet[1]]
    print("Headers:", [h for h in headers if h is not None])
    
    # Print first 5 rows
    row_count = 0
    for r in range(2, sheet.max_row + 1):
        row_values = [sheet.cell(r, c).value for c in range(1, len(headers) + 1)]
        if any(v is not None for v in row_values):
            row_count += 1
            if row_count <= 5:
                print(f"Row {row_count}:", row_values)
    print(f"Total Rows: {row_count}")
